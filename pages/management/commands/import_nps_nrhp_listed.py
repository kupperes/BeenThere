from datetime import datetime
from io import BytesIO
from urllib.request import urlopen

from django.core.management.base import BaseCommand
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from pages.models import HistoricSite, ImportRun, SourceFeed


DEFAULT_NPS_NRHP_LISTED_URL = (
    'https://www.nps.gov/subjects/nationalregister/upload/'
    'national-register-listed_20250624.xlsx'
)


def clean_cell(value):
    if value is None:
        return ''
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def normalize_header(value):
    return clean_cell(value).lower()


def parse_excel_date(value):
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    return clean_cell(value)


def build_summary(record):
    parts = ['Listed in the National Register of Historic Places']
    if record['city'] or record['state']:
        location = ', '.join(part for part in [record['city'], record['state']] if part)
        parts.append(f'for {location}')
    if record['category_of_property']:
        parts.append(f"as a {record['category_of_property'].lower()}")
    if record['listed_date']:
        parts.append(f"on {record['listed_date']}")
    return ' '.join(parts) + '.'


class Command(BaseCommand):
    help = 'Import the official NPS NRHP listed-properties spreadsheet.'

    def add_arguments(self, parser):
        parser.add_argument(
            '--source-url',
            default=DEFAULT_NPS_NRHP_LISTED_URL,
            help='Override the remote XLSX source URL.',
        )
        parser.add_argument(
            '--source-file',
            help='Read the XLSX spreadsheet from a local file instead of downloading it.',
        )

    def handle(self, *args, **options):
        source_feed, _ = SourceFeed.objects.get_or_create(
            slug='nps-nrhp-listed',
            defaults={
                'name': 'NPS NRHP Listed Properties',
                'jurisdiction_level': 'national',
                'jurisdiction_name': 'United States',
                'source_type': 'xlsx',
                'homepage_url': 'https://www.nps.gov/subjects/nationalregister/data-downloads.htm',
                'download_url': options['source_url'],
                'license': 'Public federal source data',
                'refresh_strategy': 'Manual import command using NPS listed-properties spreadsheet',
                'is_active': True,
            },
        )

        import_run = ImportRun.objects.create(
            source_feed=source_feed,
            status='running',
            notes='Importing NPS NRHP listed properties spreadsheet.',
        )

        try:
            workbook = self._load_workbook(options['source_file'], options['source_url'])
            worksheet = workbook.active
            headers = [normalize_header(cell.value) for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
            rows = [self._row_to_record(headers, row) for row in worksheet.iter_rows(min_row=2, values_only=True)]
            rows = [row for row in rows if row['property_name'] and row['reference_number']]

            created_count = 0
            updated_count = 0

            with transaction.atomic():
                for record in rows:
                    defaults = {
                        'name': record['property_name'],
                        'summary': build_summary(record),
                        'description': clean_cell(record['area_of_significance']),
                        'latitude': None,
                        'longitude': None,
                        'address': record['street_and_number'],
                        'city': record['city'],
                        'state': record['state'],
                        'category': self._map_category(record['category_of_property']),
                        'designation': 'National Register of Historic Places',
                        'era': '',
                        'source_feed': source_feed,
                        'source_name': source_feed.name,
                        'wikipedia_url': '',
                        'reference_url': 'https://www.nps.gov/subjects/nationalregister/database-research.htm',
                        'image_url': '',
                        'is_verified': True,
                    }
                    _, created = HistoricSite.objects.update_or_create(
                        source_feed=source_feed,
                        source_id=record['reference_number'],
                        defaults=defaults,
                    )
                    if created:
                        created_count += 1
                    else:
                        updated_count += 1

            import_run.imported_count = len(rows)
            import_run.created_count = created_count
            import_run.updated_count = updated_count
            import_run.status = 'completed'
            import_run.finished_at = timezone.now()
            import_run.notes = (
                f'Imported from approved NPS source: '
                f'{options["source_file"] or options["source_url"]}'
            )
            import_run.save(update_fields=[
                'imported_count',
                'created_count',
                'updated_count',
                'status',
                'finished_at',
                'notes',
            ])

            self.stdout.write(
                self.style.SUCCESS(
                    f'Imported {len(rows)} NPS NRHP listed properties '
                    f'({created_count} created, {updated_count} updated).'
                )
            )
        except Exception as exc:
            import_run.status = 'failed'
            import_run.failed_count = 1
            import_run.error_log = str(exc)
            import_run.finished_at = timezone.now()
            import_run.save(update_fields=['status', 'failed_count', 'error_log', 'finished_at'])
            raise

    def _load_workbook(self, source_file, source_url):
        if source_file:
            return load_workbook(filename=source_file, read_only=True)

        with urlopen(source_url) as response:
            return load_workbook(filename=BytesIO(response.read()), read_only=True)

    def _row_to_record(self, headers, row):
        values = {header: row[index] if index < len(row) else None for index, header in enumerate(headers)}
        return {
            'reference_number': clean_cell(values.get('ref#')),
            'property_name': clean_cell(values.get('property name')),
            'state': clean_cell(values.get('state')),
            'county': clean_cell(values.get('county')),
            'city': clean_cell(values.get('city')),
            'street_and_number': clean_cell(values.get('street & number')),
            'area_of_significance': clean_cell(values.get('area of significance')),
            'category_of_property': clean_cell(values.get('category of property')),
            'listed_date': parse_excel_date(values.get('listed date')),
        }

    def _map_category(self, category_of_property):
        mapping = {
            'building': 'landmark',
            'district': 'district',
            'site': 'marker',
            'structure': 'landmark',
            'object': 'memorial',
        }
        return mapping.get(category_of_property.lower(), 'other') if category_of_property else 'other'
